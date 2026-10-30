"""Complete local data preparation; source content never supplies instructions.

The pilot is kept intact. This module extends its provenance conventions and
reference loader, with an explicit full-source registry and versioned schema.
"""
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sqlite3
import sys
from tempfile import TemporaryDirectory

import openpyxl
from openpyxl.utils.datetime import from_excel

from .pilot import (PilotError, assert_header, blank, connect_database, department,
                    insert, integer, issue, load_references, protected_hashes,
                    raw_json, read_source, stable_id, text)
from .sources import SOURCES
from .redaction import Redactor, VERSION as REDACTION_VERSION

VERSION = '2.0.0'
MODULE = Path(__file__).parent
REF_SOURCES = tuple(s for s in SOURCES if s.kind in {'dealerships','business'})
MODEL_FIELDS = {'event_id','text','department','text_sha256','redaction_version','input_snapshot','privacy_status'}


@dataclass
class Observation:
    source_id: str
    profile: dict
    sheet: str
    row: int
    end_row: int
    values: dict
    segment: dict
    kind: str
    epoch: datetime | None = None

    @property
    def observation_id(self):
        return stable_id('observation',self.source_id,self.sheet,self.row)

    def get(self,role):
        return self.values.get(self.segment.get('columns',{}).get(role))


def registry():
    value=json.loads((MODULE/'preparation_profiles.json').read_text(encoding='utf-8'))
    if value['registry_version']!=2:raise PilotError('Unsupported preparation registry version')
    return value['sources']


def trimmed(values):
    values=list(values)
    while values and values[-1] is None:values.pop()
    return values


def extract(root,profile):
    """All populated rows are data, registered headers or hashed context."""
    path=root/profile['path']
    try:payload=path.read_bytes()
    except OSError as exc:raise PilotError(f"Missing source: {profile['path']}") from exc
    digest=hashlib.sha256(payload).hexdigest()
    source_id=stable_id('source',profile['path'],digest)
    metadata={'sheets':[],'notes':[],'context_rows':[],'blank_rows':[]}
    observations=[]
    workbook=None
    epoch=None
    if path.suffix=='.csv':
        try:
            reader=csv.reader(io.StringIO(payload.decode(profile['encoding']),newline=''),delimiter=profile['delimiter'],strict=True)
            rows={};ends={};previous=0
            for values in reader:
                first=previous+1;previous=reader.line_num
                rows[first]=values;ends[first]=previous
            surfaces={'CSV':(rows,ends)}
            metadata.update(encoding=profile['encoding'],delimiter=profile['delimiter'])
        except (UnicodeError,csv.Error) as exc:raise PilotError(f"{profile['path']}: invalid CSV encoding/quoting") from exc
    else:
        try:workbook=openpyxl.load_workbook(io.BytesIO(payload),read_only=False,data_only=False)
        except Exception as exc:raise PilotError(f"{profile['path']}: cannot decode workbook") from exc
        epoch=workbook.epoch;metadata['excel_epoch']=epoch.isoformat()
        surfaces={}
        for sheet in workbook:
            rows={row[0].row:[cell.value for cell in row] for row in sheet}
            surfaces[sheet.title]=(rows,{row:row for row in rows})
            metadata['sheets'].append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,
                'state':sheet.sheet_state,'merged_ranges':[str(v) for v in sheet.merged_cells.ranges],
                'hidden_rows':[i for i,d in sheet.row_dimensions.items() if d.hidden],
                'hidden_columns':[i for i,d in sheet.column_dimensions.items() if d.hidden]})
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f' or cell.data_type=='e':
                        workbook.close();raise PilotError(f"{profile['path']} :: {sheet.title} :: row {cell.row}: unexpected formula/error cell")
                    if cell.comment:metadata['notes'].append({'sheet':sheet.title,'cell':cell.coordinate,'text':cell.comment.text,'author':cell.comment.author})
    try:
        if set(surfaces)!={s['name'] for s in profile['sheets']}:
            raise PilotError(f"{profile['path']}: unregistered or missing worksheet")
        for sheet in profile['sheets']:
            name=sheet['name'];rows,ends=surfaces[name]
            context={c['row']:c for c in sheet['context']};covered=set()
            for row,c in context.items():
                if row not in rows or stable_id(trimmed(rows[row]))!=c['sha256']:
                    raise PilotError(f"{profile['path']} :: {name} :: row {row}: context changed; review profile")
                metadata['context_rows'].append({'sheet':name,'row':row,'role':c['role'],'values':rows[row]});covered.add(row)
            for segment in sheet['segments']:
                hr=segment['header_row']
                assert_header(rows.get(hr,[]),segment['headers'],f"{profile['path']} :: {name} :: row {hr}")
                covered.add(hr);metadata['context_rows'].append({'sheet':name,'row':hr,'role':'header','values':rows[hr]})
                count=0
                for row,values in rows.items():
                    if row in context or not segment['first_row']<=row<=segment['last_row']:continue
                    covered.add(row)
                    if not values or all(blank(v) for v in values):
                        metadata['blank_rows'].append({'sheet':name,'row':row,'end_row':ends[row]});continue
                    n=len(segment['headers'])
                    if any(not blank(v) for v in values[n:]) or (path.suffix=='.csv' and len(values)!=n):
                        raise PilotError(f"{profile['path']} :: {name} :: row {row}: unexpected field count")
                    if sum(isinstance(v,str) and v in segment['headers'] for v in values)>=2:
                        raise PilotError(f"{profile['path']} :: {name} :: row {row}: unregistered header")
                    kind={'individual':'feedback','external_review':'external_review','aggregate':'rating_aggregate'}[sheet['kind']]
                    observations.append(Observation(source_id,profile,name,row,ends[row],dict(zip(segment['headers'],values)),segment,kind,epoch));count+=1
                if count!=segment['expected_records']:
                    raise PilotError(f"{profile['path']} :: {name} :: header {hr}: source count differs from independent control")
            for row,values in rows.items():
                if row not in covered:
                    if any(not blank(v) for v in values):
                        raise PilotError(f"{profile['path']} :: {name} :: row {row}: unaccounted populated row")
                    metadata['blank_rows'].append({'sheet':name,'row':row,'end_row':ends[row]})
            if path.suffix=='.csv':metadata['sheets'].append({'name':'CSV','physical_lines':max(ends.values(),default=0)})
        if 'Noosa feedback cards' in profile['path'] and not any(n['sheet']=='Cards' and n['cell']=='C1' for n in metadata['notes']):
            raise PilotError('Noosa required collection note missing')
        note_hashes=[{'sheet':n['sheet'],'cell':n['cell'],'sha256':stable_id(n['text'],n['author'])} for n in metadata['notes']]
        if note_hashes!=profile.get('notes',[]):
            raise PilotError(f"{profile['path']}: worksheet comments changed; review profile")
    finally:
        if workbook:workbook.close()
    return source_id,digest,metadata,observations


def interpreted_date(value,segment,epoch):
    """Return date/month/precision/basis/candidate; never invent a day."""
    fmt=segment['date_format']
    if fmt=='relative':return None,None,'relative','relative_without_capture_date',None
    if blank(value):return None,None,'unknown','missing',None
    if fmt=='%d/%m/%Y' and isinstance(value,str):
        match=re.fullmatch(r'(early|mid|wk[1-5]) ([A-Z][a-z]{2})',value.strip())
        uncertain=re.fullmatch(r'([A-Z][a-z]{2})\?',value.strip())
        if match or uncertain:
            month=datetime.strptime((match[2] if match else uncertain[1]),'%b').month
            start,end=segment['period_start'],segment['period_end']
            candidates=[f'{year}-{month:02}' for year in range(int(start[:4]),int(end[:4])+1)
                        if start[:7]<=f'{year}-{month:02}'<end[:7]]
            if len(candidates)==1:
                if uncertain:return None,None,'unknown','uncertain_month',candidates[0]
                return None,candidates[0],'month','month_text_with_year_inferred_from_source_window',None
    try:
        if fmt=='excel_datetime':
            if not isinstance(value,(date,datetime)):raise ValueError()
            parsed=value.date() if isinstance(value,datetime) else value
        elif fmt=='excel_serial':
            if isinstance(value,bool):raise ValueError()
            if isinstance(value,(date,datetime)):parsed=value.date() if isinstance(value,datetime) else value
            elif isinstance(value,(int,float)):
                result=from_excel(value,epoch=epoch)
                parsed=result.date() if isinstance(result,datetime) else result
                if not isinstance(parsed,date):raise ValueError()
            else:raise ValueError()
        else:
            if not isinstance(value,str):raise ValueError()
            parsed=datetime.strptime(value,fmt).date()
        return parsed.isoformat(),parsed.strftime('%Y-%m'),'day','declared_source_format',None
    except (ValueError,TypeError,OverflowError):return None,None,'unknown','invalid',None


def native_rating(value,segment):
    if blank(value):return None
    kind=segment['rating_kind']
    if kind=='satisfaction_labels':
        return {'very dissatisfied':1,'dissatisfied':2,'neutral':3,'satisfied':4,'very satisfied':5}.get(str(value).strip().casefold())
    if kind=='three_labels':return {'bad':1,'ok':2,'good':3}.get(str(value).strip().casefold())
    if kind=='stars':
        s=str(value).strip()
        return s.count('★') if re.fullmatch('★{1,5}☆{0,4}',s) and len(s)==5 else None
    try:
        n=integer(value)
        return n if 1<=n<=segment['scale_max'] else None
    except ValueError:return None


def normalise(observation,closed_dates):
    o=observation;s=o.segment;p=o.profile;problems=[]
    dealer=p.get('dealer');candidate=p.get('dealer_candidate');dealer_status='source_profile'
    if 'North QLD' in p['path']:
        dealer={'rockhampton':'QLD07','rocky':'QLD07','rok':'QLD07','townsville':'QLD08','tsv':'QLD08'}.get(str(o.get('site')).strip().casefold())
        dealer_status='reviewed_source_alias' if dealer else 'unresolved_alias'
    elif p.get('dealer_status'):dealer_status=p['dealer_status']
    if dealer is None:problems.append(('UNRESOLVED_DEALER','Site alias lacks a confirmed reference mapping','Retained; excluded from dealer/business comparisons'))
    if dealer_status=='assumed_mapping':
        # Decision D3: workbook title, the only dealership without another file, Hunter region, and comment brands matching only NSW05.
        problems.append(('ASSUMED_DEALER_MAPPING','"Hunter Motors" alias assumed to be NSW05 Kestrel Newcastle (sole dealer without another file; Hunter region; comment brands Altura/Kodiak/Voltra match only NSW05)','Mapped with assumed status; alias retained; exclude for sensitivity'))
    raw_dept=s.get('department_constant') or o.get('department')
    dept,subtype=department(raw_dept)
    if str(raw_dept).strip().casefold()=='parts & accessories':dept,subtype='Parts','Parts & Accessories'
    sales_subtypes={'sales - new':'New Cars','new':'New Cars','sales - used':'Used Cars','used':'Used Cars'}
    if str(raw_dept).strip().casefold() in sales_subtypes:
        dept,subtype='Sales',sales_subtypes[str(raw_dept).strip().casefold()]
    dept_candidate=None;dept_status='reviewed_alias' if dept else 'unresolved'
    if 'bendigo' in p['path']:
        # Decision D4: comment keywords align with each code and the Service share matches the group pattern.
        dept_candidate={'s':'Sales','sv':'Service','p':'Parts'}.get(str(raw_dept).casefold())
        dept,subtype,dept_status=dept_candidate,None,'assumed_code' if dept_candidate else 'unresolved'
        if dept:problems.append(('ASSUMED_DEPARTMENT_CODE','Bendigo code S/SV/P assumed to be Sales/Service/Parts (comment keywords align with each code; Service share matches the group)','Mapped with assumed status; raw code retained; exclude for sensitivity'))
    if dept is None and o.kind!='external_review':problems.append(('UNRESOLVED_DEPARTMENT','Department needs an explicit mapping decision','Retained; excluded from department/business comparisons'))
    day,month,precision,basis,candidate_month=interpreted_date(o.get('date'),s,o.epoch)
    time_ok=bool(month and s['period_start'][:7]<=month<s['period_end'][:7])
    if day:time_ok=s['period_start']<=day<s['period_end']
    if precision=='month':problems.append(('APPROXIMATE_MONTH','Month known; year inferred from the single source reporting window','Month-level use only; no exact day invented'))
    if not time_ok:problems.append(('DATE_UNUSABLE_FOR_PERIOD','Date is missing, uncertain, relative or outside the source window','Retained; excluded from period-based comparisons'))
    if dealer in closed_dates and ((day and day>closed_dates[dealer]) or (month and month>closed_dates[dealer][:7])):
        time_ok=False;problems.append(('AFTER_DEALERSHIP_CLOSURE','Recorded date is after closure','Original date retained; excluded from period-based comparisons'))
    rating=native_rating(o.get('rating'),s);direction=s['direction']
    if rating is None:problems.append(('MISSING_RATING' if blank(o.get('rating')) else 'INVALID_RATING','Native rating is missing or invalid','Text can remain usable; excluded from rating distributions'))
    comment=text(o.get('comment'));has_text=not blank(comment)
    if not has_text and rating is None:problems.append(('NO_ANALYTICAL_CONTENT','Neither a valid rating nor comment is present','Retained for reconciliation; no analytical contribution'))
    reference=text(o.get('reference'))
    return {
        'event_id':stable_id('event',o.observation_id),'canonical_observation_id':o.observation_id,
        'population':'external_review' if o.kind=='external_review' else 'individual',
        'dealer_code':dealer,'dealer_candidate':candidate,'dealer_status':dealer_status,
        'department_raw':text(raw_dept),'department':dept,'department_candidate':dept_candidate,'department_subtype':subtype,'department_status':dept_status,
        'feedback_date_raw':text(o.get('date')),'feedback_date':day,'feedback_month':month,'date_candidate_month':candidate_month,
        'date_precision':precision,'date_basis':basis,'time_eligible':int(time_ok),'day_eligible':int(time_ok and precision=='day'),
        'business_join_eligible':int(time_ok and dealer is not None and dept is not None and o.kind=='feedback' and (rating is not None or has_text)),
        'rating_raw':text(o.get('rating')),'rating_value':rating,'rating_scale_max':s['scale_max'],'rating_kind':s['rating_kind'],
        'rating_direction':direction,'rating_oriented':(s['scale_max']+1-rating if direction=='low_is_good' else rating) if rating is not None and direction!='unspecified' else None,
        'rating_valid':int(rating is not None),'comment_raw':comment,'has_text':int(has_text),'text_eligible':int(has_text and o.kind=='feedback'),
        'channel_raw':text(o.get('channel')) or ('copied_google_reviews' if o.kind=='external_review' else None),
        'native_reference':reference,'reference_namespace':f"{dealer or 'unresolved'}:feedback" if reference else None,
    },problems


def resolve(observations):
    """Return canonical links/exclusions using only known, evidenced copy rules."""
    feedback=[o for o in observations if o.kind=='feedback']
    links={};excluded={};byfile=defaultdict(list)
    for o in feedback:byfile[o.profile['path']].append(o)
    old='feedback/Chermside Q1 2026 FINAL.xlsx';new='feedback/Chermside Q1 2026 FINAL v2.xlsx'
    candidates=defaultdict(list)
    for o in byfile[new]:candidates[raw_json(o.values)].append(o)
    for o in byfile[old]:
        matches=candidates[raw_json(o.values)]
        if len(matches)!=1:raise PilotError('Chermside version overlap is not a unique exact match')
        links[o.observation_id]=(matches[0].observation_id,'chermside_exact_version_overlap')
    pen=byfile['feedback/Penrith Feedback.xlsx']
    july=defaultdict(list)
    for o in pen:
        if o.sheet=='Jul 26':july[raw_json(o.values)].append(o)
    for o in pen:
        if o.sheet=='Jun 26':
            matches=july[raw_json(o.values)]
            if len(matches)!=1:raise PilotError('Penrith copied-month evidence changed')
            links[o.observation_id]=(matches[0].observation_id,'penrith_exact_july_copy')
    ess=byfile['feedback/ESSENDON - feedback register.xlsx']
    def escalation_key(o):return raw_json([o.get(k) for k in ('date','rating','comment','staff')])
    service=defaultdict(list)
    for o in ess:
        if o.sheet=='Service':service[escalation_key(o)].append(o)
    for o in ess:
        if o.sheet=='Sales' and str(o.get('action')).startswith('ESCALATED from Service'):
            matches=service[escalation_key(o)]
            if len(matches)!=1:raise PilotError('Essendon escalation has no unique Service original')
            links[o.observation_id]=(matches[0].observation_id,'essendon_routed_service_copy')
    seen={}
    for o in byfile['feedback/Springwood feedback.xlsx']:
        sig=raw_json(o.values)
        if sig in seen:links[o.observation_id]=(seen[sig],'springwood_complete_row_repeat')
        else:seen[sig]=o.observation_id
    expected={23:'test test',167:'testing new form',192:'asdf',196:'test - ignore'}
    for o in byfile['feedback/Southport Feedback Tracker.xlsx']:
        if o.sheet=='Feedback' and o.row in expected:
            if str(o.get('comment')).strip().casefold()!=expected[o.row]:raise PilotError('Southport test-row evidence changed')
            excluded[o.observation_id]='identified_test_record'
    counts=Counter(rule for _,rule in links.values())
    if counts!=Counter(chermside_exact_version_overlap=112,penrith_exact_july_copy=43,essendon_routed_service_copy=13,springwood_complete_row_repeat=12) or len(excluded)!=4:
        raise PilotError('Evidenced-copy/test controls differ; investigate rather than targeting counts')
    return links,excluded


def load_observations(db,observations,links,excluded,metadata_by_source):
    closed=dict(db.execute("SELECT dealer_code,closed_date FROM dealerships WHERE closed_date IS NOT NULL"))
    for o in observations:
        disposition='aggregate' if o.kind=='rating_aggregate' else 'excluded' if o.observation_id in excluded else 'pending'
        insert(db,'source_observations',{
            'observation_id':o.observation_id,'source_id':o.source_id,'sheet':o.sheet,'source_row':o.row,'source_end_row':o.end_row,
            'layout_version':f"{o.sheet}:header-{o.segment['header_row']}",'kind':o.kind,'raw_fields_json':raw_json(o.values),
            'event_id':None,'disposition':disposition,'resolution_rule':excluded.get(o.observation_id,'aggregate_source' if o.kind=='rating_aggregate' else 'canonical_source_row')})
    for o in observations:
        if o.kind=='rating_aggregate':
            counts=[integer(o.values[str(i)]) for i in range(1,6)]
            if min(counts)<0 or sum(counts)!=integer(o.values['Grand Total']):raise PilotError('Aggregate row totals do not reconcile')
            dept,_=department(o.values['Department'])
            if dept is None:raise PilotError('Unknown aggregate department')
            for bucket,count in enumerate(counts,1):
                insert(db,'rating_aggregates',{'observation_id':o.observation_id,'dealer_code':o.profile['dealer'],
                    'month':datetime.strptime(o.values['Month'],'%b-%y').strftime('%Y-%m'),'department':dept,
                    'rating_bucket':bucket,'response_count':count,'rating_direction':'unspecified'})
            continue
        if o.observation_id in excluded:
            issue(db,o.source_id,'IDENTIFIED_TEST_ROW','Known test entry excluded from analytical events','Original observation retained',record=o,severity='info')
            continue
        if o.observation_id in links:continue
        event,problems=normalise(o,closed)
        insert(db,'feedback_events',event)
        db.execute("UPDATE source_observations SET event_id=?,disposition='retained' WHERE observation_id=?",(event['event_id'],o.observation_id))
        for rule,message,action in problems:issue(db,o.source_id,rule,message,action,record=o)
        if o.kind=='external_review':issue(db,o.source_id,'EXTERNAL_REVIEW_UNCERTAIN','Copied review has relative timing and unverified overlap','Separate population; not in primary feedback rates or model input',record=o)
    for source_id,meta in metadata_by_source.items():
        for context in meta.get('context_rows',[]):
            if context['role']=='control_total' and context['values'][0]=='Grand Total':
                expected=[integer(v) for v in context['values'][2:8]]
                observed=[db.execute('SELECT SUM(response_count) FROM rating_aggregates a JOIN source_observations o USING(observation_id) WHERE o.source_id=? AND rating_bucket=?',(source_id,i)).fetchone()[0] for i in range(1,6)]
                if observed!=expected[:5] or sum(observed)!=expected[5]:raise PilotError('Aggregate source grand total differs from bucket totals')
            if context['role']=='footer_or_note' and str(context['values'][0]).startswith('Total '):
                # Frankston supplies independent monthly count/rounded-mean controls.
                month=datetime.strptime(context['values'][0],'Total %b %Y').strftime('%Y-%m')
                count,mean=db.execute('''SELECT COUNT(*),AVG(rating_value) FROM feedback_events e
                  JOIN source_observations o ON o.observation_id=e.canonical_observation_id
                  WHERE o.source_id=? AND feedback_month=?''',(source_id,month)).fetchone()
                supplied=int(str(context['values'][3]).split()[0])
                if count!=supplied or round(mean,2)!=float(context['values'][2]):
                    raise PilotError('Monthly source subtotal count/mean does not reconcile')
        if any(n['cell']=='C1' and n['sheet']=='Cards' for n in meta.get('notes',[])):
            issue(db,source_id,'SELECTIVE_COLLECTION_NOTE','A source note describes selective feedback collection','Context retained; not a representative customer sample',sheet='Cards',row=1)
    for observation_id,(canonical,rule) in links.items():
        db.execute("UPDATE source_observations SET event_id=?,disposition='duplicate',resolution_rule=? WHERE observation_id=?",
                   (stable_id('event',canonical),rule,observation_id))
    # Keep low-information collisions, rather than deleting them using a weak key.
    candidates=defaultdict(list)
    for o in observations:
        if o.kind=='feedback' and o.observation_id not in links and o.observation_id not in excluded:
            candidates[(o.profile['path'],raw_json(o.values))].append(o)
    for group in candidates.values():
        if len(group)<2:continue
        group_id=stable_id('possible-duplicate',sorted(o.observation_id for o in group))
        for o in group:
            db.execute('UPDATE feedback_events SET duplicate_candidate_group=? WHERE canonical_observation_id=?',(group_id,o.observation_id))
            issue(db,o.source_id,'POSSIBLE_DUPLICATE','Identical retained observations have insufficient copy evidence','Retained; use duplicate group for sensitivity and evaluation splits',record=o)
    for profile in registry():
        source_id=next(o.source_id for o in observations if o.profile['path']==profile['path'])
        for sheet in profile['sheets']:
            for segment in sheet['segments']:
                if segment.get('direction','unspecified')=='unspecified':
                    issue(db,source_id,'RATING_DIRECTION_UNSPECIFIED','No supplied good/poor direction is established','Native distributions only; no negative-rating bands inferred',sheet=sheet['name'],row=segment['header_row'])


def make_redactor(observations):
    known=defaultdict(set)
    for o in observations:
        for role,kind in [('name','name'),('staff','name'),('phone','phone'),('email','email'),('rego','rego'),('reference','id'),('customer_id','id')]:
            value=o.get(role)
            if blank(value):continue
            value=str(value).strip();known[kind].add(value)
    return Redactor(known)


def apply_redaction(db,observations):
    redactor=make_redactor(observations)
    lookup={o.observation_id:o for o in observations}
    for row in db.execute('SELECT event_id,canonical_observation_id,comment_raw,has_text,text_eligible FROM feedback_events').fetchall():
        if not row['has_text']:
            db.execute("UPDATE feedback_events SET privacy_status='not_applicable',redaction_version=? WHERE event_id=?",(REDACTION_VERSION,row['event_id']));continue
        redacted,counts,reasons=redactor.redact(row['comment_raw'])
        status='external_channel_held' if not row['text_eligible'] else 'withheld' if reasons else 'automated_pass'
        db.execute('UPDATE feedback_events SET redacted_text=?,redaction_version=?,redaction_count=?,redaction_tags_json=?,privacy_reasons_json=?,privacy_status=?,model_eligible=? WHERE event_id=?',
                   (redacted,REDACTION_VERSION,sum(counts.values()),raw_json(counts),raw_json(reasons),status,int(status=='automated_pass'),row['event_id']))
        if reasons and row['text_eligible']:
            o=lookup[row['canonical_observation_id']]
            issue(db,o.source_id,'PRIVACY_REVIEW_REQUIRED','Automated redaction has unresolved identifier/name signals','Withheld from model input; review redacted text locally',record=o)
    return redactor


def validate_prepared(db,observations,profiles,originals_unchanged):
    checks=[]
    def expect(name,actual,expected):checks.append({'name':name,'actual':actual,'expected':expected,'passed':actual==expected})
    def scalar(sql,params=()):return db.execute(sql,params).fetchone()[0]
    expect('Registered files',scalar('SELECT COUNT(*) FROM source_files'),54)
    for p in profiles:
        for sheet in p['sheets']:
            for segment in sheet['segments']:
                actual=scalar('''SELECT COUNT(*) FROM source_observations o JOIN source_files f USING(source_id)
                  WHERE f.path=? AND sheet=? AND layout_version=?''',(p['path'],sheet['name'],f"{sheet['name']}:header-{segment['header_row']}"))
                expect(f"Source segment: {p['path']} :: {sheet['name']} :: {segment['header_row']}",actual,segment['expected_records'])
    for name,sql,n in [
        ('Individual observations',"SELECT COUNT(*) FROM source_observations WHERE kind='feedback'",7257),
        ('Retained individual events',"SELECT COUNT(*) FROM feedback_events WHERE population='individual'",7073),
        ('External copied reviews',"SELECT COUNT(*) FROM feedback_events WHERE population='external_review'",8),
        ('Exact copy links',"SELECT COUNT(*) FROM source_observations WHERE disposition='duplicate'",180),
        ('Identified tests retained as exclusions',"SELECT COUNT(*) FROM source_observations WHERE disposition='excluded'",4),
        ('All observations accounted',"SELECT COUNT(*) FROM source_observations",len(observations)),
        ('No pending dispositions',"SELECT COUNT(*) FROM source_observations WHERE disposition='pending'",0),
        ('Canonical links agree',"SELECT COUNT(*) FROM feedback_events e JOIN source_observations o ON e.canonical_observation_id=o.observation_id AND e.event_id=o.event_id AND o.disposition='retained'",7081),
        ('No excluded/aggregate events',"SELECT COUNT(*) FROM source_observations WHERE disposition IN ('excluded','aggregate') AND event_id IS NOT NULL",0),
        ('Separate aggregate source rows',"SELECT COUNT(*) FROM source_observations WHERE kind='rating_aggregate'",6),
        ('Weighted aggregate responses',"SELECT SUM(response_count) FROM rating_aggregates",71),
        ('Dealership reference',"SELECT COUNT(*) FROM dealerships",24),
        ('Business department-month scopes',"SELECT COUNT(*) FROM business_monthly",864),
        ('Unique dealer-month business keys',"SELECT COUNT(*) FROM (SELECT dealer_code,month FROM business_monthly GROUP BY dealer_code,month)",288),
        ('Source revenue in cents',"SELECT SUM(revenue_cents) FROM business_monthly",83868599000),
        ('Business join retains all scopes',"SELECT COUNT(*) FROM business_context",864),
        ('Business join preserves revenue',"SELECT SUM(revenue_cents) FROM business_context",83868599000),
        ('Missing feedback stays unavailable',"SELECT COUNT(*) FROM business_context WHERE feedback_status='no_eligible_feedback' AND event_count IS NOT NULL",0),
        ('Unknown rating direction never oriented',"SELECT COUNT(*) FROM feedback_events WHERE rating_direction='unspecified' AND rating_oriented IS NOT NULL",0),
        ('External reviews not model inputs',"SELECT COUNT(*) FROM feedback_events WHERE population='external_review' AND model_eligible=1",0),
        ('Withheld comments not model inputs',"SELECT COUNT(*) FROM feedback_events WHERE privacy_status='withheld' AND model_eligible=1",0),
        ('Approximate dates have no invented day',"SELECT COUNT(*) FROM feedback_events WHERE date_precision='month' AND feedback_date IS NOT NULL",0),
        ('Assumed Hunter dealer mapping visible',"SELECT COUNT(*) FROM feedback_events WHERE dealer_code='NSW05' AND dealer_candidate='NSW05' AND dealer_status='assumed_mapping'",303),
        ('Assumed Hunter mapping findings',"SELECT COUNT(*) FROM quality_issues WHERE rule_id='ASSUMED_DEALER_MAPPING'",303),
        ('No unresolved Hunter dealer aliases',"SELECT COUNT(*) FROM feedback_events WHERE dealer_code IS NULL AND dealer_candidate='NSW05'",0),
        ('Assumed Bendigo department codes visible',"SELECT COUNT(*) FROM feedback_events WHERE dealer_code='VIC08' AND department_status='assumed_code' AND department=department_candidate AND department_raw IN ('S','SV','P')",144),
        ('Assumed Bendigo code findings',"SELECT COUNT(*) FROM quality_issues WHERE rule_id='ASSUMED_DEPARTMENT_CODE'",144),
        ('No unresolved Bendigo departments',"SELECT COUNT(*) FROM feedback_events WHERE dealer_code='VIC08' AND department IS NULL",0),
        ('Incorrect Shepparton dates excluded from periods',"SELECT COUNT(*) FROM feedback_events WHERE dealer_code='VIC09' AND feedback_date>='2026-10-01' AND time_eligible=0",6),
        ('Missing rating and comment retained',"SELECT COUNT(*) FROM feedback_events WHERE rating_valid=0 AND has_text=0",2),
        ('Empty responses contribute no analytical counts',"SELECT COUNT(*) FROM feedback_events WHERE rating_valid=0 AND has_text=0 AND business_join_eligible=1",0),
    ]:expect(name,scalar(sql),n)
    expect('SQLite integrity',scalar('PRAGMA integrity_check'),'ok')
    expect('Foreign-key violations',len(db.execute('PRAGMA foreign_key_check').fetchall()),0)
    expect('Foreign-key enforcement',scalar('PRAGMA foreign_keys'),1)
    expect('Original inputs unchanged',originals_unchanged,True)
    return checks


def report_data(db,checks,run_id):
    counts={
      'individual_observations':db.execute("SELECT COUNT(*) FROM source_observations WHERE kind='feedback'").fetchone()[0],
      'individual_events':db.execute("SELECT COUNT(*) FROM feedback_events WHERE population='individual'").fetchone()[0],
      'duplicate_observations':db.execute("SELECT COUNT(*) FROM source_observations WHERE disposition='duplicate'").fetchone()[0],
      'excluded_tests':db.execute("SELECT COUNT(*) FROM source_observations WHERE disposition='excluded'").fetchone()[0],
      'external_reviews':db.execute("SELECT COUNT(*) FROM feedback_events WHERE population='external_review'").fetchone()[0],
      'aggregate_responses':db.execute('SELECT SUM(response_count) FROM rating_aggregates').fetchone()[0]}
    sources=[dict(r) for r in db.execute('''SELECT f.path,o.kind,COUNT(*) observations,
      SUM(o.disposition='retained') retained,SUM(o.disposition='duplicate') copies,
      SUM(o.disposition='excluded') excluded,SUM(o.disposition='aggregate') aggregate_rows
      FROM source_files f JOIN source_observations o USING(source_id)
      GROUP BY f.path,o.kind ORDER BY f.path,o.kind''')]
    issues=[dict(r) for r in db.execute('SELECT rule_id,severity,COUNT(*) findings FROM quality_issues GROUP BY rule_id,severity ORDER BY rule_id')]
    privacy=[dict(r) for r in db.execute('SELECT population,privacy_status,COUNT(*) events FROM feedback_events GROUP BY population,privacy_status ORDER BY population,privacy_status')]
    eligible=dict(db.execute("""SELECT 'valid_native_ratings',SUM(rating_valid) FROM feedback_events WHERE population='individual'
      UNION ALL SELECT 'usable_period',SUM(time_eligible) FROM feedback_events WHERE population='individual'
      UNION ALL SELECT 'usable_exact_day',SUM(day_eligible) FROM feedback_events WHERE population='individual'
      UNION ALL SELECT 'business_join',SUM(business_join_eligible) FROM feedback_events WHERE population='individual'
      UNION ALL SELECT 'primary_text',SUM(text_eligible) FROM feedback_events
      UNION ALL SELECT 'model_input',SUM(model_eligible) FROM feedback_events"""))
    reasons=Counter(reason for r in db.execute("SELECT privacy_reasons_json FROM feedback_events WHERE privacy_status='withheld'") for reason in json.loads(r[0]))
    coverage=[dict(r) for r in db.execute('''SELECT dealer_code,department,feedback_month,date_precision,COUNT(*) events,
      SUM(time_eligible) time_eligible,SUM(business_join_eligible) business_join_eligible,SUM(model_eligible) model_eligible
      FROM feedback_events WHERE population='individual' GROUP BY dealer_code,department,feedback_month,date_precision
      ORDER BY dealer_code,department,feedback_month,date_precision''')]
    return {'run_id':run_id,'schema_version':2,'parser_version':VERSION,'redaction_version':REDACTION_VERSION,
            'counts':counts,'sources':sources,'issues':issues,'privacy':privacy,'eligibility':eligible,
            'privacy_review_reasons':dict(sorted(reasons.items())),'coverage':coverage,'checks':checks,
            'limitations':['Automated redaction plus targeted review is not guaranteed anonymity or independent human privacy certification.',
              'Hunter Motors -> NSW05 Newcastle (D3) and Bendigo S/SV/P -> Sales/Service/Parts (D4) are evidence-based assumptions, flagged by status and quality finding with raw values retained; rerun results without them as a sensitivity check.',
              'Native scales and unknown directions remain separate; no negative bands or cross-scale means are inferred.',
              'Copied Google reviews have unknown capture dates/overlap and remain a separate population.',
              'Redaction rules-3 replaces residual names it can locate with [NAME] and withholds only remaining contact, identifier or name signals; unknown single first names without a person cue can still pass.',
              'Missing feedback, collection bias, ambiguous duplicates and uncertain dates remain visible.',
              'No AI labels, human reference annotations, causal effects or financial uplift have been produced.']}


def validate_model_input(db,path,input_snapshot,redactor):
    """Validate the actual serialized boundary, including source alignment."""
    expected={r['event_id']:dict(r) for r in db.execute('''SELECT event_id,redacted_text,department,redaction_version
      FROM feedback_events WHERE model_eligible=1''')}
    seen=set();shape_ok=True;alignment_ok=True;privacy_ok=True
    try:
        with path.open(encoding='utf-8') as stream:
            for line in stream:
                value=json.loads(line)
                if not isinstance(value,dict) or set(value)!=MODEL_FIELDS:
                    shape_ok=False;continue
                key=value['event_id'];row=expected.get(key)
                if row is None or key in seen:
                    alignment_ok=False;continue
                seen.add(key)
                if (value['text']!=row['redacted_text'] or value['department']!=row['department']
                    or value['redaction_version']!=row['redaction_version'] or value['input_snapshot']!=input_snapshot
                    or value['privacy_status']!='automated_pass'
                    or value['text_sha256']!=hashlib.sha256(row['redacted_text'].encode()).hexdigest()):
                    alignment_ok=False
                _,counts,reasons=redactor.redact(value['text'])
                if counts or reasons:privacy_ok=False
    except (ValueError,TypeError):shape_ok=False
    values=[('Model input has only allowlisted fields',shape_ok),
            ('Model input exactly matches eligible events and versions',alignment_ok and seen==set(expected)),
            ('Model input has no remaining detected identifiers/name signals',privacy_ok)]
    return [{'name':name,'actual':passed,'expected':True,'passed':passed} for name,passed in values]


def write_outputs(db,directory,checks,run_id,input_snapshot,redactor):
    with (directory/'model_input.jsonl').open('w',encoding='utf-8') as stream:
        for row in db.execute('SELECT event_id,redacted_text,department,redaction_version FROM feedback_events WHERE model_eligible=1 ORDER BY event_id'):
            stream.write(json.dumps({'event_id':row['event_id'],'text':row['redacted_text'],'department':row['department'],
                'text_sha256':hashlib.sha256(row['redacted_text'].encode()).hexdigest(),'redaction_version':row['redaction_version'],
                'input_snapshot':input_snapshot,'privacy_status':'automated_pass'},ensure_ascii=False)+'\n')
    checks.extend(validate_model_input(db,directory/'model_input.jsonl',input_snapshot,redactor))
    if not all(c['passed'] for c in checks):raise PilotError('Model-input boundary validation failed')
    report=report_data(db,checks,run_id)
    (directory/'quality_report.json').write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    lines=['# Stage 2 — Complete data preparation','',
      'Status: built and validated; user checkpoint required before the AI pilot.','',
      '**7,257 individual observations → 7,073 retained individual events; 180 copy links and four test exclusions.**',
      'Eight copied Google reviews and 71 aggregate responses remain separate. All 52 feedback files and both references are registered.','',
      f"Checks: **{sum(c['passed'] for c in checks)}/{len(checks)}** passed. Run: `{run_id}`.",'',
      '## Privacy and model handoff','',
      '| Population | Privacy status | Events |','|---|---|---:|']
    lines += [f"| {r['population']} | {r['privacy_status']} | {r['events']} |" for r in report['privacy']]
    eligible=report['eligibility']
    lines += ['',f"Model-input coverage: **{eligible['model_input']:,}/{eligible['primary_text']:,} primary text comments ({eligible['model_input']/eligible['primary_text']:.1%})**. Missing text and external reviews are separate from this denominator."]
    lines += ['', 'Only automated-pass eligible individual comments appear in `model_input.jsonl`. This is an allowlisted redacted export; it contains no raw tables, customer IDs, contact columns, source names or numeric ratings.',
      'Withheld redacted comments and their issue codes are available locally in the private database for review. Automated-pass status is not independent privacy approval.','',
      '## Decisions and limitations','']+[f'- {v}' for v in report['limitations']]
    month_count=db.execute("SELECT COUNT(*) FROM feedback_events WHERE date_precision='month'").fetchone()[0]
    uncertain_count=db.execute("SELECT COUNT(*) FROM feedback_events WHERE date_basis='uncertain_month'").fetchone()[0]
    lines += [f'- {month_count} approximate month entries infer their year from the single reporting window; no day is invented. {uncertain_count} `Aug?` entries remain uncertain candidate months.',
      '- Invalid/missing dates and ratings affect their own eligibility; valid comments or ratings are preserved independently.',
      '- Every copy retains its original source fields, including Essendon action histories. Ambiguous complete-row collisions are flagged, not deleted.',
      '- The 24 dealerships and 288 dealer/month records reshape into 864 department/month scopes, retaining AUD 838,685,990. Missing eligible feedback stays NULL.','',
      '## Quality findings','', '| Rule | Severity | Findings |','|---|---|---:|']
    lines += [f"| {r['rule_id']} | {r['severity']} | {r['findings']} |" for r in report['issues']]
    lines += ['', '## Source reconciliation','', '| Source | Kind | Observations | Retained | Copies | Excluded | Aggregate rows |','|---|---|---:|---:|---:|---:|---:|']
    lines += [f"| {Path(r['path']).name} | {r['kind']} | {r['observations']} | {r['retained']} | {r['copies']} | {r['excluded']} | {r['aggregate_rows']} |" for r in report['sources']]
    lines += ['', 'Database: `prepared.sqlite`. Detailed checks and coverage: `quality_report.json`. Input/schema/parser/redaction versions and source hashes: database metadata and `manifest.json`.',
      'Query `event_counts`, `native_rating_distribution`, `business_context`, `event_source_history` and `quality_issues` locally. The database contains raw personal information and must remain private.','']
    (directory/'summary.md').write_text('\n'.join(lines),encoding='utf-8')
    return report


def code_fingerprint():
    files=['preparation.py','preparation_profiles.json','prepared_schema.sql','redaction.py','pilot.py','sources.py']
    return stable_id([(name,hashlib.sha256((MODULE/name).read_bytes()).hexdigest()) for name in files])


def run_preparation(root,output_dir=None):
    root=Path(root).resolve();output_dir=Path(output_dir or root/'outputs/prepared').resolve()
    if output_dir==root or any(output_dir.is_relative_to((root/folder).resolve()) for folder in ('feedback','reference','archive','kestrel','tests','docs','outputs/pilot')):
        raise PilotError('Output directory cannot replace the project or supplied/source directories')
    profiles=registry()
    expected={p['path'] for p in profiles}|{s.path for s in REF_SOURCES}
    try:
        actual={str(p.relative_to(root)) for folder in ('feedback','reference') for p in (root/folder).iterdir() if p.name!='.DS_Store'}
    except OSError as exc:raise PilotError('Cannot read the required source directories') from exc
    if actual!=expected:raise PilotError('Source inventory has missing or unregistered files')
    before=protected_hashes(root)
    refs=[read_source(root,s) for s in REF_SOURCES]
    extracted=[(p,*extract(root,p)) for p in profiles]
    observations=[o for p,sid,digest,meta,obs in extracted for o in obs]
    links,excluded=resolve(observations)
    snapshot=stable_id(sorted([(e.source.path,e.sha256) for e in refs]+[(p['path'],digest) for p,sid,digest,meta,obs in extracted]))
    code_hash=code_fingerprint()
    runtime={'python':sys.version.split()[0],'sqlite':sqlite3.sqlite_version,'openpyxl':openpyxl.__version__}
    run_id=stable_id('prepared',snapshot,code_hash,VERSION,REDACTION_VERSION,runtime)
    output_dir.mkdir(parents=True,exist_ok=True,mode=0o700)
    current=output_dir/'current'
    if current.exists() and not current.is_symlink():raise PilotError('Current output pointer is not a managed symlink')
    with TemporaryDirectory(prefix='.prepare-',dir=output_dir) as temp:
        directory=Path(temp);db=connect_database(directory/'prepared.sqlite')
        try:
            db.executescript((MODULE/'prepared_schema.sql').read_text())
            for e in refs:
                insert(db,'source_files',{'source_id':e.source_id,'path':e.source.path,'sha256':e.sha256,'kind':e.source.kind,
                    'profile_json':raw_json(e.source.profile()),'metadata_json':raw_json(e.metadata)})
            for p,sid,digest,meta,obs in extracted:
                insert(db,'source_files',{'source_id':sid,'path':p['path'],'sha256':digest,'kind':'feedback',
                    'profile_json':raw_json({'registry_version':2,**p}),'metadata_json':raw_json(meta)})
            load_references(db,refs)
            load_observations(db,observations,links,excluded,{sid:meta for p,sid,digest,meta,obs in extracted})
            redactor=apply_redaction(db,observations)
            checks=validate_prepared(db,observations,profiles,before==protected_hashes(root))
            failed=[c['name'] for c in checks if not c['passed']]
            if failed:raise PilotError('Reconciliation failed: '+'; '.join(failed))
            metadata={'schema_version':'2','parser_version':VERSION,'registry_version':'2','redaction_version':REDACTION_VERSION,
                'input_snapshot':snapshot,'run_id':run_id,'code_sha256':code_hash,'python_version':sys.version.split()[0],
                'sqlite_version':sqlite3.sqlite_version,'openpyxl_version':openpyxl.__version__,
                'protected_inputs_sha256':raw_json(before),'validation_checks':raw_json(checks)}
            report=write_outputs(db,directory,checks,run_id,snapshot,redactor)
            metadata['validation_checks']=raw_json(checks)
            for key,value in metadata.items():insert(db,'run_metadata',{'key':key,'value':value})
            db.commit()
        except sqlite3.DatabaseError as exc:raise PilotError('Prepared database constraint/integrity failure') from exc
        finally:db.close()
        artifacts={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(directory.iterdir())}
        manifest={'run_id':run_id,'input_snapshot':snapshot,'code_sha256':code_hash,'schema_version':2,
            'redaction_version':REDACTION_VERSION,'source_sha256':before,'artifacts':artifacts,'runtime':runtime,
            'model_input_status':'automatically_checked; user review checkpoint pending','model_calls':0}
        (directory/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
        for p in directory.iterdir():p.chmod(0o600)
        runs=output_dir/'runs';runs.mkdir(exist_ok=True,mode=0o700);target=runs/run_id
        if target.exists():
            if {p.name for p in target.iterdir()}!={p.name for p in directory.iterdir()} or any(not (target/p.name).is_file() or (target/p.name).read_bytes()!=p.read_bytes() for p in directory.iterdir()):
                raise PilotError('Existing immutable run differs; investigate repeatability')
        else:
            # A complete run directory is installed first; one symlink swap
            # publishes the database, redacted input and report together.
            os.rename(directory,target)
        pointer=output_dir/(directory.name+'-link')
        try:
            pointer.symlink_to(Path('runs')/run_id)
            os.replace(pointer,current)
        finally:
            if pointer.is_symlink():pointer.unlink()
    return {'database':current/'prepared.sqlite','report':current/'summary.md','model_input':current/'model_input.jsonl',
            'quality_report':current/'quality_report.json','manifest':current/'manifest.json','checks':checks,'run_id':run_id,'counts':report['counts']}

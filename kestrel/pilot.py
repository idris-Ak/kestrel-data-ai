"""Deterministic local import, reconciliation and reporting for stage one."""

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import sqlite3
import sys
from tempfile import TemporaryDirectory

import openpyxl

from . import __version__
from .sources import (
    AGGREGATE_HEADER, BUSINESS_HEADER, DEALER_HEADER, NEWER_CHERMSIDE,
    OLDER_CHERMSIDE, SOURCES, EXPECTED_RECORDS, Segment, Source,
)


class PilotError(Exception):
    """An actionable import/reconciliation failure, without customer values."""


def raw_json(value):
    def preserve_nonfinite(obj):
        if isinstance(obj, float) and not math.isfinite(obj):
            return {"type": "float", "value": str(obj)}
        if isinstance(obj, dict):
            return {k: preserve_nonfinite(v) for k, v in obj.items()}
        if isinstance(obj, (list, tuple)):
            return [preserve_nonfinite(v) for v in obj]
        return obj

    def encode(obj):
        if isinstance(obj, (datetime, date)):
            return {"type": type(obj).__name__, "value": obj.isoformat()}
        raise TypeError(f"Unsupported source value type: {type(obj).__name__}")
    return json.dumps(preserve_nonfinite(value), sort_keys=True, ensure_ascii=False, default=encode,
                      separators=(",", ":"), allow_nan=False)


def stable_id(*parts):
    return hashlib.sha256(raw_json(parts).encode()).hexdigest()


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def text(value):
    return None if value is None else value.isoformat() if isinstance(value, (date, datetime)) else str(value)


def integer(value):
    if isinstance(value, bool) or blank(value):
        raise ValueError("Expected an integer")
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError("Expected a finite integer")
        result = int(number)
        if not -(2**63) <= result < 2**63:
            raise ValueError("Integer outside SQLite range")
        return result
    except InvalidOperation as exc:
        raise ValueError("Expected an integer") from exc


def money_cents(value):
    try:
        amount = Decimal(str(value)) * 100
        result = integer(amount)
        if result < 0:
            raise ValueError("Negative revenue")
        return result
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("Expected nonnegative revenue with exact cents") from exc


DEPARTMENTS = {
    "sales": ("Sales", None), "new cars": ("Sales", "New Cars"),
    "used cars": ("Sales", "Used Cars"), "service": ("Service", None),
    "service dept": ("Service", None), "svc": ("Service", None),
    "workshop": ("Service", None), "parts": ("Parts", None),
    "p&a": ("Parts", "P&A"), "parts counter": ("Parts", "Parts counter"),
}


def department(value):
    return DEPARTMENTS.get(str(value).strip().casefold(), (None, None))


@dataclass
class Record:
    source: Source
    source_id: str
    sheet: str
    row: int
    end_row: int
    values: dict
    segment: Segment | None

    @property
    def observation_id(self):
        return stable_id("observation", self.source_id, self.sheet, self.row)

    @property
    def layout(self):
        return f"{self.sheet}:header-{self.segment.header_row if self.segment else 3}"


@dataclass
class Extracted:
    source: Source
    source_id: str
    sha256: str
    metadata: dict
    records: list[Record]


def assert_header(actual, expected, location):
    actual = list(actual)
    while actual and actual[-1] is None:
        actual.pop()
    if tuple(actual) != tuple(expected):
        raise PilotError(f"{location}: unrecognised header; review the source profile")


def read_source(root: Path, source: Source) -> Extracted:
    path = root / source.path
    try:
        payload = path.read_bytes()
    except OSError as exc:
        raise PilotError(f"Cannot read required input: {source.path}") from exc
    digest = hashlib.sha256(payload).hexdigest()
    source_id = stable_id("source", source.path, digest)
    metadata = {"context_rows": [], "blank_rows": [], "sheets": [], "notes": []}
    records = []
    if path.suffix == ".csv":
        try:
            stream = io.StringIO(payload.decode(source.encoding), newline="")
            reader = csv.reader(stream, delimiter=source.delimiter, strict=True)
            header = next(reader)
            expected = DEALER_HEADER if source.kind == "dealerships" else BUSINESS_HEADER if source.kind == "business" else source.segments[0].headers
            assert_header(header, expected, f"{source.path} :: CSV :: row 1")
            metadata["context_rows"].append({"sheet": "CSV", "row": 1, "role": "header", "values": header})
            metadata["encoding"] = source.encoding
            metadata["delimiter"] = source.delimiter
            previous_line = reader.line_num
            for values in reader:
                first, last = previous_line + 1, reader.line_num
                previous_line = last
                if not values or all(blank(v) for v in values):
                    metadata["blank_rows"].append({"sheet": "CSV", "row": first, "end_row": last})
                    continue
                if len(values) != len(header):
                    raise PilotError(f"{source.path} :: CSV :: row {first}: unexpected field count")
                if sum(v in header for v in values) >= 2:
                    raise PilotError(f"{source.path} :: CSV :: row {first}: unregistered header inside data")
                records.append(Record(source, source_id, "CSV", first, last,
                                      dict(zip(header, values)), source.segments[0] if source.segments else None))
            metadata["sheets"].append({"name": "CSV", "physical_lines": previous_line})
        except (UnicodeError, csv.Error, StopIteration) as exc:
            raise PilotError(f"{source.path}: invalid CSV encoding, quoting or missing header") from exc
    else:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=False, data_only=False)
        except Exception as exc:
            raise PilotError(f"{source.path}: cannot decode workbook") from exc
        try:
            if set(workbook.sheetnames) != set(source.expected_sheets):
                raise PilotError(f"{source.path}: unexpected worksheet set")
            metadata["excel_epoch"] = workbook.epoch.isoformat()
            for sheet in workbook:
                metadata["sheets"].append({"name": sheet.title, "rows": sheet.max_row,
                                           "columns": sheet.max_column, "state": sheet.sheet_state,
                                           "merged_ranges": [str(r) for r in sheet.merged_cells.ranges]})
                for row in sheet:
                    for cell in row:
                        if cell.comment:
                            metadata["notes"].append({"sheet": sheet.title, "cell": cell.coordinate,
                                                       "text": cell.comment.text, "author": cell.comment.author})
            for sheet, cell in source.required_notes:
                if workbook[sheet][cell].comment is None:
                    raise PilotError(f"{source.path} :: {sheet}!{cell}: required context note missing")
            context = {(s, r): role for s, r, role in source.context_rows}
            covered = set(context)
            segments = source.segments
            if source.kind == "aggregate":
                assert_header([c.value for c in workbook["Summary"][3]], AGGREGATE_HEADER,
                              f"{source.path} :: Summary :: row 3")
                covered.add(("Summary", 3))
                metadata["context_rows"].append({"sheet": "Summary", "row": 3, "role": "header", "values": list(AGGREGATE_HEADER)})
                for row in range(4, 10):
                    values = [c.value for c in workbook["Summary"][row]]
                    if any(not blank(v) for v in values[len(AGGREGATE_HEADER):]):
                        raise PilotError(f"{source.path} :: Summary :: row {row}: unexpected extra columns")
                    records.append(Record(source, source_id, "Summary", row, row,
                                          dict(zip(AGGREGATE_HEADER, values)), None))
                    covered.add(("Summary", row))
            for segment in segments:
                sheet = workbook[segment.sheet]
                assert_header([c.value for c in sheet[segment.header_row]], segment.headers,
                              f"{source.path} :: {segment.sheet} :: row {segment.header_row}")
                covered.add((segment.sheet, segment.header_row))
                metadata["context_rows"].append({"sheet": segment.sheet, "row": segment.header_row,
                                                  "role": "header", "values": list(segment.headers)})
                for row_no in range(segment.first_row, (segment.last_row or sheet.max_row) + 1):
                    values = [c.value for c in sheet[row_no]]
                    covered.add((segment.sheet, row_no))
                    if all(blank(v) for v in values):
                        metadata["blank_rows"].append({"sheet": segment.sheet, "row": row_no})
                        continue
                    if any(not blank(v) for v in values[len(segment.headers):]):
                        raise PilotError(f"{source.path} :: {segment.sheet} :: row {row_no}: unexpected extra columns")
                    if sum(isinstance(v, str) and v in segment.headers for v in values) >= 2:
                        raise PilotError(f"{source.path} :: {segment.sheet} :: row {row_no}: unregistered header inside data")
                    records.append(Record(source, source_id, segment.sheet, row_no, row_no,
                                          dict(zip(segment.headers, values)), segment))
            for (sheet, row), role in context.items():
                values = [c.value for c in workbook[sheet][row]]
                if role == "blank" and any(not blank(v) for v in values):
                    raise PilotError(f"{source.path} :: {sheet} :: row {row}: expected blank context row")
                metadata["context_rows"].append({"sheet": sheet, "row": row, "role": role, "values": values})
            for sheet in workbook:
                for row in sheet:
                    if any(not blank(c.value) for c in row) and (sheet.title, row[0].row) not in covered:
                        raise PilotError(f"{source.path} :: {sheet.title} :: row {row[0].row}: unaccounted populated row")
            if source.path.endswith("Geelong Feedback.xlsx"):
                if "1 = Excellent, 5 = Poor" not in str(workbook["2025"]["A2"].value) or "1 = Poor" not in str(workbook["2026"]["A2"].value) or "5 = Excellent" not in str(workbook["2026"]["A2"].value):
                    raise PilotError(f"{source.path}: rating legend changed; review conversion")
            if "Ringwood" in source.path and workbook["Feedback"]["A36"].value != "*** NEW TEMPLATE FROM FEB 2026 ***":
                raise PilotError(f"{source.path}: layout-change notice changed")
        finally:
            workbook.close()
    return Extracted(source, source_id, digest, metadata, records)


def connect_database(path):
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys=ON")
    if connection.execute("PRAGMA foreign_keys").fetchone()[0] != 1:
        raise PilotError("SQLite foreign-key enforcement unavailable")
    return connection


def insert(db, table, values):
    # Identifiers are developer-owned table/field names. Source values are bound.
    columns = ",".join(values)
    db.execute(f"INSERT INTO {table} ({columns}) VALUES ({','.join('?' for _ in values)})", tuple(values.values()))


def issue(db, source_id, rule, message, disposition, *, record=None, sheet=None, row=None, severity="warning"):
    observation = record.observation_id if record else None
    sheet = record.sheet if record else sheet
    row = record.row if record else row
    insert(db, "quality_issues", {"issue_id": stable_id("issue", source_id, observation, sheet, row, rule),
        "source_id": source_id, "observation_id": observation, "sheet": sheet, "source_row": row,
        "severity": severity, "rule_id": rule, "message": message, "disposition": disposition})


def parse_date(value, segment):
    if blank(value):
        raise ValueError("Missing date")
    if segment.date_format == "excel_datetime":
        if not isinstance(value, (date, datetime)):
            raise ValueError("Expected native Excel date")
        return value.date() if isinstance(value, datetime) else value
    if not isinstance(value, str):
        raise ValueError("Expected source-format date text")
    return datetime.strptime(value, segment.date_format).date()


def normalise_event(record):
    """Return an event and field issues. No inferred sentiment or missing values."""
    segment, values = record.segment, record.values
    problems = []
    dept, subtype = department(values[segment.department_column])
    if dept is None:
        problems.append(("UNKNOWN_DEPARTMENT", "Department is unresolved", "Retained; excluded from department/business joins"))
    value_date = None
    time_eligible = False
    try:
        value_date = parse_date(values[segment.date_column], segment).isoformat()
        time_eligible = segment.period_start <= value_date < segment.period_end
        if not time_eligible:
            problems.append(("DATE_OUTSIDE_SOURCE_PERIOD", "Date conflicts with the source period", "Retained; excluded from period-based results"))
    except (ValueError, OverflowError):
        problems.append(("MISSING_DATE" if blank(values[segment.date_column]) else "INVALID_DATE",
                         "Date is missing or does not match the source format", "Retained; excluded from period-based results"))
    rating = None
    try:
        rating = integer(values[segment.rating_column])
        if not 1 <= rating <= segment.scale_max:
            raise ValueError("Rating out of native scale")
    except ValueError:
        rating = None
        problems.append(("MISSING_RATING" if blank(values[segment.rating_column]) else "INVALID_RATING",
                         "Rating is missing or outside the native scale", "Retained; excluded from rating distributions"))
    oriented = rating if segment.direction == "high_is_good" else segment.scale_max + 1 - rating if rating is not None and segment.direction == "low_is_good" else None
    comment = text(values[segment.comment_column])
    reference = text(values.get(segment.reference_column)) if segment.reference_column else None
    return {
        "event_id": stable_id("event", record.observation_id), "canonical_observation_id": record.observation_id,
        "dealer_code": record.source.dealer, "department_raw": text(values[segment.department_column]),
        "department": dept, "department_subtype": subtype,
        "feedback_date_raw": text(values[segment.date_column]), "feedback_date": value_date,
        "date_precision": "day" if value_date else "unknown", "time_eligible": int(time_eligible),
        "rating_raw": text(values[segment.rating_column]), "rating_value": rating,
        "rating_scale_max": segment.scale_max, "rating_direction": segment.direction,
        "rating_oriented": oriented, "rating_valid": int(rating is not None),
        "comment_raw": comment, "has_text": int(not blank(comment)),
        "native_reference": reference, "reference_namespace": f"{record.source.dealer}:feedback" if reference else None,
    }, problems


def load_references(db, extracted):
    for source in extracted:
        if source.source.kind not in {"dealerships", "business"}:
            continue
        for record in source.records:
            values = record.values
            location = f"{source.source.path} :: row {record.row}"
            try:
                if source.source.kind == "dealerships":
                    closed = values["closed_date"] or None
                    if closed:
                        date.fromisoformat(closed)
                    insert(db, "dealerships", {**{k: values[k] for k in DEALER_HEADER if k != "closed_date"},
                           "closed_date": closed, "source_id": source.source_id,
                           "source_row": record.row, "raw_fields_json": raw_json(values)})
                else:
                    month = values["month"]
                    if datetime.strptime(month, "%Y-%m").strftime("%Y-%m") != month:
                        raise ValueError("Invalid month")
                    for dept, activity, amount in (
                        ("Sales", "vehicles_sold", "sales_revenue_aud"),
                        ("Service", "service_jobs", "service_revenue_aud"),
                        ("Parts", "parts_transactions", "parts_revenue_aud"),
                    ):
                        count = integer(values[activity])
                        if count < 0:
                            raise ValueError("Negative activity")
                        insert(db, "business_monthly", {"dealer_code": values["dealer_code"], "month": month,
                               "department": dept, "activity_count": count, "revenue_cents": money_cents(values[amount]),
                               "source_id": source.source_id, "source_row": record.row, "raw_fields_json": raw_json(values)})
            except (ValueError, sqlite3.IntegrityError) as exc:
                raise PilotError(f"{location}: invalid or nonunique reference record") from exc


def load_feedback(db, extracted):
    feedback = [e for e in extracted if e.source.kind in {"feedback", "aggregate"}]
    for source in feedback:
        for record in source.records:
            insert(db, "source_observations", {"observation_id": record.observation_id, "source_id": source.source_id,
                "sheet": record.sheet, "source_row": record.row, "source_end_row": record.end_row,
                "layout_version": record.layout, "kind": "feedback" if source.source.kind == "feedback" else "rating_aggregate",
                "raw_fields_json": raw_json(record.values), "event_id": None,
                "disposition": "pending" if source.source.kind == "feedback" else "aggregate", "resolution_rule": None})
    for source in feedback:
        if source.source.kind == "aggregate":
            load_aggregate(db, source)
            continue
        if source.source.path == OLDER_CHERMSIDE:
            continue
        for record in source.records:
            event, problems = normalise_event(record)
            insert(db, "feedback_events", event)
            db.execute("UPDATE source_observations SET event_id=?, disposition='retained', resolution_rule='canonical_source_row' WHERE observation_id=?",
                       (event["event_id"], record.observation_id))
            for rule, message, disposition in problems:
                issue(db, source.source_id, rule, message, disposition, record=record)
        for segment in source.source.segments:
            if segment.direction == "unspecified":
                issue(db, source.source_id, "RATING_DIRECTION_UNSPECIFIED", "Numeric bounds are known; a good/poor direction is not stated",
                      "Native distributions only; no oriented score or negative threshold inferred",
                      sheet=segment.sheet, row=segment.header_row)
        if source.source.required_notes:
            issue(db, source.source_id, "SELECTIVE_COLLECTION_NOTE", "A source note describes selectively distributing feedback cards",
                  "Context retained; do not interpret the sample as representative", sheet="Cards", row=1)
    newer = next(e for e in feedback if e.source.path == NEWER_CHERMSIDE)
    older = next(e for e in feedback if e.source.path == OLDER_CHERMSIDE)
    matches = defaultdict(list)
    for record in newer.records:
        matches[raw_json(record.values)].append(record)
    for record in older.records:
        candidates = matches[raw_json(record.values)]
        if len(candidates) != 1:
            raise PilotError(f"{OLDER_CHERMSIDE} :: row {record.row}: version overlap is not a unique full-row match")
        canonical = candidates[0]
        db.execute("UPDATE source_observations SET event_id=?,disposition='duplicate',resolution_rule='chermside_exact_version_overlap' WHERE observation_id=?",
                   (stable_id("event", canonical.observation_id), record.observation_id))
    issue(db, older.source_id, "SUPERSEDED_VERSION", "Older source observations exactly match rows in the newer version",
          "Source rows retained and linked to newer canonical events", severity="info")


def load_aggregate(db, source):
    totals = [0] * 5
    for record in source.records:
        values = record.values
        try:
            month = datetime.strptime(values["Month"], "%b-%y").strftime("%Y-%m")
            dept, _ = department(values["Department"])
            counts = [integer(values[str(i)]) for i in range(1, 6)]
            if dept is None or not "2025-10" <= month <= "2025-12" or min(counts) < 0 or sum(counts) != integer(values["Grand Total"]):
                raise ValueError("Aggregate mismatch")
            for i, count in enumerate(counts, 1):
                insert(db, "rating_aggregates", {"observation_id": record.observation_id,
                       "dealer_code": source.source.dealer, "month": month, "department": dept,
                       "rating_bucket": i, "response_count": count, "rating_direction": "unspecified"})
                totals[i-1] += count
        except (ValueError, TypeError) as exc:
            raise PilotError(f"{source.source.path} :: row {record.row}: aggregate counts do not reconcile") from exc
    control = next(c for c in source.metadata["context_rows"] if c["role"] == "control_total")["values"]
    try:
        if any(not blank(v) for v in control[8:]) or control[0] != "Grand Total" or totals != [integer(v) for v in control[2:7]] or sum(totals) != integer(control[7]):
            raise ValueError("Control mismatch")
    except (ValueError, IndexError) as exc:
        raise PilotError(f"{source.source.path}: grand total does not reconcile") from exc
    issue(db, source.source_id, "AGGREGATE_ONLY", "This source supplies rating frequencies without individual comments",
          "Kept as weighted rating buckets; no feedback events invented", severity="info")
    issue(db, source.source_id, "RATING_DIRECTION_UNSPECIFIED", "Bucket numbers do not include a good/poor direction legend",
          "Preserved as native counts; no negative category inferred", sheet="Summary", row=3)


def protected_hashes(root):
    paths = [root / "docs" / "BRIEF.md", root / "PS_Data_AI_Assignment.pdf"]
    paths.extend(p for folder in (root / "feedback", root / "reference") for p in folder.iterdir() if p.is_file())
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(paths) if p.is_file()}


def scalar(db, query, params=()):
    return db.execute(query, params).fetchone()[0]


def validate(db, originals_unchanged):
    checks = []
    def expect(name, actual, expected):
        checks.append({"name": name, "actual": actual, "expected": expected, "passed": actual == expected})
    expect("Selected source files", scalar(db, "SELECT COUNT(*) FROM source_files"), 10)
    for source in SOURCES:
        if source.kind == "dealerships":
            count = scalar(db, "SELECT COUNT(*) FROM dealerships")
        elif source.kind == "business":
            count = scalar(db, "SELECT COUNT(DISTINCT source_row) FROM business_monthly")
        else:
            count = scalar(db, "SELECT COUNT(*) FROM source_observations o JOIN source_files f USING(source_id) WHERE f.path=?", (source.path,))
        expect(f"Source records: {Path(source.path).name}", count, EXPECTED_RECORDS[source.path])
    expect("Individual source observations", scalar(db, "SELECT COUNT(*) FROM source_observations WHERE kind='feedback'"), 844)
    expect("Resolved individual events", scalar(db, "SELECT COUNT(*) FROM feedback_events"), 732)
    expect("Linked version copies", scalar(db, "SELECT COUNT(*) FROM source_observations WHERE disposition='duplicate'"), 112)
    expect("Unassigned observations", scalar(db, "SELECT COUNT(*) FROM source_observations WHERE disposition='pending' OR (kind='feedback' AND event_id IS NULL)"), 0)
    expect("Canonical event/source links", scalar(db, "SELECT COUNT(*) FROM feedback_events e JOIN source_observations o ON e.canonical_observation_id=o.observation_id AND e.event_id=o.event_id AND o.disposition='retained'"), 732)
    expect("Geelong reversed ratings", scalar(db, "SELECT COUNT(*) FROM feedback_events e JOIN source_observations o ON o.observation_id=e.canonical_observation_id WHERE e.dealer_code='VIC06' AND o.sheet='2025' AND rating_oriented=6-rating_value"), 68)
    expect("Geelong newer ratings preserved", scalar(db, "SELECT COUNT(*) FROM feedback_events e JOIN source_observations o ON o.observation_id=e.canonical_observation_id WHERE e.dealer_code='VIC06' AND o.sheet='2026' AND rating_oriented=rating_value"), 236)
    for scale, expected in ((5, 33), (10, 79)):
        expect(f"Ringwood native {scale}-point observations", scalar(db, "SELECT COUNT(*) FROM feedback_events WHERE dealer_code='VIC02' AND rating_scale_max=?", (scale,)), expected)
    expect("Liverpool dates in declared quarter", scalar(db, "SELECT COUNT(*) FROM feedback_events WHERE dealer_code='NSW04' AND feedback_date>='2026-07-01' AND feedback_date<'2026-10-01' AND time_eligible=1"), 85)
    expect("Noosa collection note retained", scalar(db, "SELECT COUNT(*) FROM source_files WHERE path='feedback/Noosa feedback cards.xlsx' AND EXISTS(SELECT 1 FROM json_each(metadata_json,'$.notes') WHERE json_extract(value,'$.sheet')='Cards' AND json_extract(value,'$.cell')='C1' AND length(json_extract(value,'$.text'))>0)"), 1)
    expect("Aggregate source rows", scalar(db, "SELECT COUNT(*) FROM source_observations WHERE kind='rating_aggregate'"), 6)
    expect("Aggregate bucket rows including zeros", scalar(db, "SELECT COUNT(*) FROM rating_aggregates"), 30)
    expect("Weighted aggregate responses", scalar(db, "SELECT SUM(response_count) FROM rating_aggregates"), 71)
    expect("Aggregate observations do not create events", scalar(db, "SELECT COUNT(*) FROM source_observations WHERE kind='rating_aggregate' AND event_id IS NOT NULL"), 0)
    expect("Dealership reference rows", scalar(db, "SELECT COUNT(*) FROM dealerships"), 24)
    expect("Original dealer/month reference keys", scalar(db, "SELECT COUNT(*) FROM (SELECT dealer_code,month FROM business_monthly GROUP BY dealer_code,month)"), 288)
    expect("Department/month business rows", scalar(db, "SELECT COUNT(*) FROM business_monthly"), 864)
    months = ["2025-10", "2025-11", "2025-12"] + [f"2026-{m:02}" for m in range(1, 10)]
    expected_scopes = {(r[0], m, d) for r in db.execute("SELECT dealer_code FROM dealerships") for m in months for d in ("Sales", "Service", "Parts")}
    observed_scopes = {tuple(r) for r in db.execute("SELECT dealer_code,month,department FROM business_monthly")}
    expect("Complete expected business scopes", observed_scopes == expected_scopes, True)
    expect("Recorded closed-site zero scopes", scalar(db, "SELECT COUNT(*) FROM business_monthly b JOIN dealerships d USING(dealer_code) WHERE d.status='Closed' AND b.month>substr(d.closed_date,1,7) AND activity_count=0 AND revenue_cents=0"), 18)
    expect("Business reference revenue cents", scalar(db, "SELECT SUM(revenue_cents) FROM business_monthly"), 83868599000)
    expect("Business-context rows retained", scalar(db, "SELECT COUNT(*) FROM business_context"), 864)
    expect("Business-context revenue cents", scalar(db, "SELECT SUM(revenue_cents) FROM business_context"), 83868599000)
    expect("Missing pilot feedback is not zero", scalar(db, "SELECT COUNT(*) FROM business_context WHERE pilot_feedback_status='no_eligible_pilot_feedback' AND event_count IS NOT NULL"), 0)
    expect("Business scopes without eligible pilot feedback exist", scalar(db, "SELECT COUNT(*)>0 FROM business_context WHERE pilot_feedback_status='no_eligible_pilot_feedback'"), 1)
    expect("Foreign-key violations", len(db.execute("PRAGMA foreign_key_check").fetchall()), 0)
    expect("Foreign-key enforcement enabled", scalar(db, "PRAGMA foreign_keys"), 1)
    expect("SQLite integrity", scalar(db, "PRAGMA integrity_check"), "ok")
    expect("Original inputs unchanged", originals_unchanged, True)
    return checks


DEMO_QUERIES = {
    "Event counts": "SELECT dealer_code, department, month, event_count,\n       valid_rating_count, text_response_count\nFROM event_counts\nORDER BY dealer_code, department, month",
    "Native rating distribution": "SELECT * FROM native_rating_distribution\nORDER BY dealer_code, department, rating_scale_max,\n         rating_direction, rating_value",
    "Business totals and feedback coverage": "SELECT pilot_feedback_status, COUNT(*) AS scopes,\n       SUM(revenue_cents) AS revenue_cents\nFROM business_context\nGROUP BY pilot_feedback_status\nORDER BY pilot_feedback_status",
}


def write_summary(db, path, checks):
    lines = ["# Stage 1 — Data foundation pilot", "", "Status: built and reconciled; awaiting user confirmation before stage two.", "",
             "Eight feedback files and both reference tables. Counts describe this pilot, not unique customers or group-wide feedback coverage.", "",
             "## Source reconciliation", "", "| Source | Individual observations | Retained events | Version copies | Aggregate source rows |",
             "|---|---:|---:|---:|---:|"]
    query = """SELECT f.path,
       SUM(CASE WHEN o.kind='feedback' THEN 1 ELSE 0 END) observations,
       SUM(CASE WHEN o.disposition='retained' THEN 1 ELSE 0 END) events,
       SUM(CASE WHEN o.disposition='duplicate' THEN 1 ELSE 0 END) duplicates,
       SUM(CASE WHEN o.kind='rating_aggregate' THEN 1 ELSE 0 END) aggregates
       FROM source_files f LEFT JOIN source_observations o USING(source_id)
       WHERE f.kind IN ('feedback','aggregate') GROUP BY f.path ORDER BY f.path"""
    for row in db.execute(query):
        lines.append(f"| {Path(row['path']).name} | {row['observations']} | {row['events']} | {row['duplicates']} | {row['aggregates']} |")
    lines += ["", "**844 individual observations → 732 events; 112 version links retained.** Separately, six aggregate rows represent 71 responses. Reference data: 24 dealerships, 288 dealer/month keys, 864 department/month rows; revenue AUD 838,685,990.",
              "", "## Verification", "", f"{sum(c['passed'] for c in checks)}/{len(checks)} reconciliation checks passed. These checks cover this data foundation, not AI accuracy or every remaining file family.", ""]
    lines += ["Checks cover source-level counts, canonical links, rating conversions/scales, Liverpool dates, Noosa context, aggregate controls, reference completeness, business joins, SQLite integrity and input preservation. Exact expected/observed results are stored as JSON under `run_metadata.validation_checks`."]
    lines += ["", "Rating/date controls: Geelong retains 68 reversed and 236 unchanged ratings; Ringwood retains 33 five-point and 79 ten-point observations; all 85 Liverpool dates fall in the declared quarter."]
    lines += ["", "## Schema", "", "| Table | One row represents |", "|---|---|",
              "| source_files | A source file, its hash, profile and worksheet context |",
              "| source_observations | A feedback or aggregate source row, with original fields and disposition |",
              "| feedback_events | A resolved individual response, with native values and eligibility flags |",
              "| rating_aggregates | A supplied rating bucket and frequency, linked to its source row |",
              "| dealerships | A reference dealership |",
              "| business_monthly | A dealership/month/department business record; money is integer cents |",
              "| quality_issues | A rule finding, source location and handling decision |",
              "| run_metadata | A version, input snapshot or validation result |"]
    lines += ["", "## Open issues and counting rules", "",
              "- Numeric scales without a direction legend remain unspecified. No negative-rating bands, NPS or cross-scale average is inferred.",
              "- Noosa's collection note is retained; it describes selective feedback collection.",
              "- Aggregate ratings remain separate weighted buckets; missing comments still permit valid ratings.",
              "- Business-context joins retain all business rows. Unavailable pilot feedback is NULL, not zero complaints.",
              "- Department aliases are explicit; New Cars, Used Cars and parts subtypes retain their original context.",
              "- The SQLite database contains raw source fields and comments locally. Redaction and AI preparation belong to stage two; this report includes no customer text or contacts.", ""]
    for row in db.execute("SELECT rule_id,severity,COUNT(*) n FROM quality_issues GROUP BY rule_id,severity ORDER BY rule_id"):
        lines.append(f"- Issue `{row['rule_id']}` ({row['severity']}): {row['n']} record(s). Details and source locations are in `quality_issues`.")
    lines += ["", "## Query examples", "", "Open `pilot.sqlite` locally. These views return counts and business amounts without customer text.", ""]
    for label, query in DEMO_QUERIES.items():
        lines += [f"**{label}**", "", "```sql", query + ";", "```", ""]
    lines += ["Source hashes, profiles, notes and structural rows: `source_files`. Original row values and event links: `source_observations`. Data snapshot and schema/runtime versions: `run_metadata`.",
              "", "Run: `uv run --locked python -m kestrel pilot`. Tests: `uv run --locked python -m unittest discover -s tests -v`.", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def run_pilot(root: Path, output_dir: Path | None = None):
    root = root.resolve()
    if not root.is_dir() or not all((root / folder).is_dir() for folder in ("feedback", "reference")):
        raise PilotError("Input root must contain the supplied feedback/ and reference/ directories")
    output_dir = (output_dir or root / "outputs/pilot").resolve()
    if output_dir == root or any(output_dir.is_relative_to((root / folder).resolve()) for folder in ("feedback", "reference", "archive")):
        raise PilotError("Output directory must not replace the project root or supplied inputs/archive")
    before = protected_hashes(root)
    extracted = [read_source(root, s) for s in SOURCES]
    output_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
    with TemporaryDirectory(prefix=".pilot-", dir=output_dir) as temporary:
        temporary = Path(temporary)
        db_path = temporary / "pilot.sqlite"
        db = connect_database(db_path)
        try:
            db.executescript(Path(__file__).with_name("schema.sql").read_text())
            for source in extracted:
                insert(db, "source_files", {"source_id": source.source_id, "path": source.source.path,
                       "sha256": source.sha256, "kind": source.source.kind,
                       "profile_json": raw_json(source.source.profile()), "metadata_json": raw_json(source.metadata)})
            for key, value in {"schema_version": "1", "parser_version": __version__, "python_version": sys.version.split()[0],
                               "sqlite_version": sqlite3.sqlite_version, "openpyxl_version": openpyxl.__version__,
                               "input_snapshot": stable_id([(e.source.path, e.sha256) for e in extracted]),
                               "protected_input_count": str(len(before))}.items():
                insert(db, "run_metadata", {"key": key, "value": value})
            load_references(db, extracted)
            load_feedback(db, extracted)
            unchanged = before == protected_hashes(root)
            checks = validate(db, unchanged)
            failed = [c["name"] for c in checks if not c["passed"]]
            if failed:
                raise PilotError("Reconciliation failed: " + "; ".join(failed))
            insert(db, "run_metadata", {"key": "validation_checks", "value": raw_json(checks)})
            insert(db, "run_metadata", {"key": "protected_inputs_sha256", "value": raw_json(before)})
            for query in DEMO_QUERIES.values():
                db.execute(query).fetchall()
            write_summary(db, temporary / "summary.md", checks)
            db.commit()
        except sqlite3.DatabaseError as exc:
            raise PilotError("Database constraint or integrity failure; investigate source/profile relationships") from exc
        finally:
            db.close()
        for name in ("pilot.sqlite", "summary.md"):
            os.chmod(temporary / name, 0o600)
            os.replace(temporary / name, output_dir / name)
    return {"database": output_dir / "pilot.sqlite", "report": output_dir / "summary.md", "checks": checks}

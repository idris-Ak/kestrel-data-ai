"""Independent source controls, failure fixtures and repeatability checks.

Only temporary outputs and in-memory workbook objects are changed by tests.
Real customer values are never printed as assertion diffs or used as fixtures.
"""

import csv
from dataclasses import replace
import hashlib
import io
import json
from pathlib import Path
import re
import sqlite3
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import openpyxl

from kestrel.pilot import (
    DEMO_QUERIES, PilotError, Record, connect_database, insert, integer,
    issue, money_cents, normalise_event, protected_hashes, raw_json,
    read_source, run_pilot,
)
from kestrel.sources import SOURCES

ROOT = Path(__file__).resolve().parents[1]
BOX = next(s for s in SOURCES if "boxhill" in s.path)
LIVERPOOL = next(s for s in SOURCES if "liverpool" in s.path)
GEELONG = next(s for s in SOURCES if "Geelong" in s.path)
NOOSA = next(s for s in SOURCES if "Noosa" in s.path)


def fixture_record(**changes):
    values = dict(zip(BOX.segments[0].headers,
                      ("000123", "2025-10-08 09:30:00", "Service", "4",
                       "Synthetic feedback marker only", "EV", "000999")))
    values.update(changes)
    return Record(BOX, "synthetic-source", "CSV", 2, 2, values, BOX.segments[0])


def snapshot(db):
    """Compare logical content, not incidental SQLite page representation."""
    tables = [r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%' ORDER BY name")]
    digest = hashlib.sha256()
    for name in tables:
        digest.update(name.encode())
        rows = sorted(raw_json(list(r)) for r in db.execute(f'SELECT * FROM "{name}"'))
        for row in rows:
            digest.update(row.encode())
    return digest.hexdigest()


class RecordRulesTests(unittest.TestCase):
    def test_missing_comment_preserves_valid_rating(self):
        event, problems = normalise_event(fixture_record(verbatim=""))
        self.assertEqual((event["rating_valid"], event["has_text"]), (1, 0))
        self.assertEqual(problems, [])

    def test_invalid_rating_preserves_raw_and_usable_text(self):
        event, problems = normalise_event(fixture_record(csat="99"))
        self.assertEqual(event["rating_raw"], "99")
        self.assertIsNone(event["rating_value"])
        self.assertEqual((event["rating_valid"], event["has_text"]), (0, 1))
        self.assertEqual(problems[0][0], "INVALID_RATING")

    def test_missing_rating_does_not_turn_into_zero(self):
        event, problems = normalise_event(fixture_record(csat=""))
        self.assertIsNone(event["rating_value"])
        self.assertEqual(problems[0][0], "MISSING_RATING")

    def test_invalid_date_is_retained_with_unknown_precision(self):
        event, problems = normalise_event(fixture_record(submitted_at="31/02/2026"))
        self.assertEqual(event["feedback_date_raw"], "31/02/2026")
        self.assertIsNone(event["feedback_date"])
        self.assertEqual((event["date_precision"], event["time_eligible"]), ("unknown", 0))
        self.assertEqual(problems[0][0], "INVALID_DATE")
        self.assertEqual(event["rating_valid"], 1)

    def test_out_of_period_date_is_not_rewritten(self):
        event, problems = normalise_event(fixture_record(submitted_at="2026-08-08 09:30:00"))
        self.assertEqual(event["feedback_date"], "2026-08-08")
        self.assertEqual(event["time_eligible"], 0)
        self.assertEqual(problems[0][0], "DATE_OUTSIDE_SOURCE_PERIOD")

    def test_source_format_rejects_a_date_only_timestamp(self):
        event, _ = normalise_event(fixture_record(submitted_at="2025-10-08"))
        self.assertIsNone(event["feedback_date"])

    def test_ambiguous_liverpool_date_uses_us_order(self):
        values = {"Date": "08/09/2026", "Department": "Service", "Customer": "Synthetic",
                  "Phone": "", "Rego": "", "Rating": "3", "Feedback": "Synthetic"}
        record = Record(LIVERPOOL, "synthetic", "CSV", 2, 2, values, LIVERPOOL.segments[0])
        event, problems = normalise_event(record)
        self.assertEqual(event["feedback_date"], "2026-08-09")
        self.assertEqual(problems, [])

    def test_unknown_department_is_not_guessed(self):
        event, problems = normalise_event(fixture_record(department="Something new"))
        self.assertIsNone(event["department"])
        self.assertEqual(event["department_raw"], "Something new")
        self.assertEqual(problems[0][0], "UNKNOWN_DEPARTMENT")

    def test_alias_and_sales_subtype_preserve_original(self):
        event, _ = normalise_event(fixture_record(department="Used Cars"))
        self.assertEqual((event["department"], event["department_subtype"], event["department_raw"]),
                         ("Sales", "Used Cars", "Used Cars"))

    def test_native_id_is_not_coerced_or_used_as_event_key(self):
        event, _ = normalise_event(fixture_record())
        self.assertEqual(event["native_reference"], "000123")
        self.assertNotEqual(event["event_id"], "000123")
        another, _ = normalise_event(replace(fixture_record(), row=3, end_row=3))
        self.assertNotEqual(event["event_id"], another["event_id"])

    def test_unspecified_direction_stays_unspecified(self):
        event, _ = normalise_event(fixture_record())
        self.assertEqual(event["rating_direction"], "unspecified")
        self.assertIsNone(event["rating_oriented"])

    def test_money_uses_exact_cents_and_rejects_bad_amounts(self):
        self.assertEqual(money_cents("0.29"), 29)
        self.assertEqual(money_cents("120.10"), 12010)
        for invalid in ("0.001", "-1.00", "NaN", "Infinity", "", "text", "1e40"):
            with self.subTest(case=invalid):
                with self.assertRaises(ValueError):
                    money_cents(invalid)

    def test_nonintegral_and_boolean_counts_are_rejected(self):
        for invalid in ("1.5", True, None, "NaN"):
            with self.subTest(case=str(invalid)):
                with self.assertRaises(ValueError):
                    integer(invalid)

    def test_nonfinite_native_rating_can_be_retained_as_invalid_source_data(self):
        record=fixture_record(csat=float('nan'))
        raw=json.loads(raw_json(record.values))
        self.assertEqual(raw['csat'],{'type':'float','value':'nan'})
        event,problems=normalise_event(record)
        self.assertEqual(event['rating_valid'],0)
        self.assertEqual(problems[0][0],'INVALID_RATING')


class ReaderFailureTests(unittest.TestCase):
    def csv_source(self, payload):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        target = root / BOX.path
        target.parent.mkdir(parents=True)
        target.write_text(payload, encoding="utf-8")
        return root

    def test_unknown_header_stops_without_echoing_private_value(self):
        root = self.csv_source("person@example.test,Unexpected\n1,2\n")
        with self.assertRaises(PilotError) as caught:
            read_source(root, BOX)
        self.assertIn("unrecognised header", str(caught.exception))
        self.assertNotIn("person@example.test", str(caught.exception))

    def test_extra_csv_fields_stop_import(self):
        root = self.csv_source(",".join(BOX.segments[0].headers) + "\n1,2,3\n")
        with self.assertRaisesRegex(PilotError, "unexpected field count"):
            read_source(root, BOX)

    def test_embedded_header_stops_csv_import(self):
        header = ",".join(BOX.segments[0].headers)
        root = self.csv_source(header + "\n" + header + "\n")
        with self.assertRaisesRegex(PilotError, "unregistered header"):
            read_source(root, BOX)

    def test_csv_preserves_strings_and_multiline_row_locations(self):
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow(BOX.segments[0].headers)
        writer.writerow(["000123", "2025-10-08 09:30:00", "Service", "4", "Synthetic\nsecond line", "EV", "000999"])
        writer.writerow(["000124", "2025-10-09 09:30:00", "Sales", "5", "Synthetic", "EV", "000998"])
        source = read_source(self.csv_source(buffer.getvalue()), BOX)
        self.assertEqual([(r.row, r.end_row) for r in source.records], [(2, 3), (4, 4)])
        self.assertEqual(source.records[0].values["dms_customer_no"], "000999")

    def test_missing_input_is_an_actionable_failure(self):
        with TemporaryDirectory() as folder:
            with self.assertRaisesRegex(PilotError, "required input"):
                read_source(Path(folder), BOX)

    def test_unknown_excel_header_stops_import_in_memory(self):
        workbook = openpyxl.load_workbook(ROOT / GEELONG.path)
        workbook["2025"]["A4"] = "Unrecognised field"
        with patch("kestrel.pilot.openpyxl.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(PilotError, "unrecognised header"):
                read_source(ROOT, GEELONG)

    def test_unregistered_sheet_stops_import_in_memory(self):
        workbook = openpyxl.load_workbook(ROOT / GEELONG.path)
        workbook.create_sheet("Unregistered")
        with patch("kestrel.pilot.openpyxl.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(PilotError, "worksheet set"):
                read_source(ROOT, GEELONG)

    def test_lost_collection_note_is_detected(self):
        workbook = openpyxl.load_workbook(ROOT / NOOSA.path)
        workbook["Cards"]["C1"].comment = None
        with patch("kestrel.pilot.openpyxl.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(PilotError, "context note missing"):
                read_source(ROOT, NOOSA)

    def test_changed_rating_legend_stops_conversion(self):
        workbook = openpyxl.load_workbook(ROOT / GEELONG.path)
        workbook["2025"]["A2"] = "Different rating policy"
        with patch("kestrel.pilot.openpyxl.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(PilotError, "rating legend changed"):
                read_source(ROOT, GEELONG)

    def test_extra_aggregate_data_column_is_not_silently_dropped(self):
        source=next(s for s in SOURCES if s.kind=='aggregate')
        workbook=openpyxl.load_workbook(ROOT/source.path)
        workbook['Summary']['I4']=999
        with patch('kestrel.pilot.openpyxl.load_workbook',return_value=workbook):
            with self.assertRaisesRegex(PilotError,'unexpected extra columns'):
                read_source(ROOT,source)


class PipelineIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = TemporaryDirectory()
        cls.before = protected_hashes(ROOT)
        cls.output = run_pilot(ROOT, Path(cls.temporary.name) / "first")
        cls.db = connect_database(cls.output["database"])

    @classmethod
    def tearDownClass(cls):
        cls.db.close()
        cls.temporary.cleanup()

    def test_all_independent_reconciliation_controls_pass(self):
        self.assertTrue(all(c["passed"] for c in self.output["checks"]))
        self.assertGreaterEqual(len(self.output["checks"]), 28)

    def test_supplied_inputs_match_archived_preimplementation_hashes(self):
        archived = json.loads((ROOT / "tests/fixtures/original_input_hashes.json").read_text())
        self.assertEqual(self.before, archived)
        self.assertEqual(protected_hashes(ROOT), self.before)

    def test_source_observations_reconcile_per_file(self):
        expected = {
            "boxhill_survey_export_2025Q4.csv": (64, 64, 0), "Geelong Feedback.xlsx": (304, 304, 0),
            "Chermside Q1 2026 FINAL.xlsx": (112, 0, 112), "Chermside Q1 2026 FINAL v2.xlsx": (118, 118, 0),
            "Ringwood feedback Q1 2026.xlsx": (112, 112, 0), "liverpool_q3_2026.csv": (85, 85, 0),
            "Noosa feedback cards.xlsx": (49, 49, 0),
        }
        rows = self.db.execute("""SELECT path,COUNT(*) n,SUM(disposition='retained') kept,SUM(disposition='duplicate') copies
            FROM source_observations JOIN source_files USING(source_id) WHERE source_observations.kind='feedback'
            GROUP BY path""").fetchall()
        actual = {Path(r["path"]).name: (r["n"], r["kept"], r["copies"]) for r in rows}
        self.assertEqual(actual, expected)

    def test_version_copies_share_event_but_retain_both_locations(self):
        row = self.db.execute("""SELECT COUNT(DISTINCT old.event_id) n
            FROM source_observations old JOIN source_observations canonical ON old.event_id=canonical.event_id
            WHERE old.disposition='duplicate' AND canonical.disposition='retained'
              AND old.observation_id<>canonical.observation_id AND old.raw_fields_json=canonical.raw_fields_json""").fetchone()
        self.assertEqual(row[0], 112)

    def test_native_ratings_and_reversals_preserve_raw_values(self):
        rows = self.db.execute("""SELECT rating_value,rating_oriented,rating_raw,o.raw_fields_json
            FROM feedback_events e JOIN source_observations o ON o.observation_id=e.canonical_observation_id
            WHERE dealer_code='VIC06' AND sheet='2025'""").fetchall()
        self.assertEqual(len(rows), 68)
        self.assertTrue(all(r["rating_oriented"] == 6-r["rating_value"] for r in rows))
        self.assertTrue(all(str(json.loads(r["raw_fields_json"])["Rating"]) == r["rating_raw"] for r in rows))

    def test_missing_comments_remain_countable_ratings(self):
        count = self.db.execute("SELECT COUNT(*) FROM feedback_events WHERE has_text=0 AND rating_valid=1").fetchone()[0]
        self.assertGreater(count, 0)

    def test_aggregate_buckets_match_source_grand_totals(self):
        counts = dict(self.db.execute("SELECT rating_bucket,SUM(response_count) FROM rating_aggregates GROUP BY rating_bucket"))
        self.assertEqual(counts, {1: 9, 2: 7, 3: 10, 4: 17, 5: 28})
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM feedback_events WHERE dealer_code='QLD06'").fetchone()[0], 0)

    def test_business_join_keeps_full_population_and_unknown_feedback(self):
        base = self.db.execute("SELECT SUM(revenue_cents),COUNT(*) FROM business_monthly").fetchone()
        joined = self.db.execute("SELECT SUM(revenue_cents),COUNT(*) FROM business_context").fetchone()
        self.assertEqual(tuple(base), (83868599000, 864))
        self.assertEqual(tuple(base), tuple(joined))
        missing = self.db.execute("SELECT COUNT(*) FROM business_context WHERE event_count IS NULL AND pilot_feedback_status='no_eligible_pilot_feedback'").fetchone()[0]
        self.assertGreater(missing, 0)

    def test_department_subtype_and_unspecified_rating_direction_are_visible(self):
        subtype = self.db.execute("SELECT COUNT(*) FROM feedback_events WHERE department='Sales' AND department_subtype='Used Cars'").fetchone()[0]
        self.assertGreater(subtype, 0)
        wrong = self.db.execute("SELECT COUNT(*) FROM feedback_events WHERE rating_direction='unspecified' AND rating_oriented IS NOT NULL").fetchone()[0]
        self.assertEqual(wrong, 0)

    def test_invalid_record_can_be_stored_with_lineage_and_issue(self):
        db = connect_database(":memory:")
        self.addCleanup(db.close)
        self.db.backup(db)
        record = replace(fixture_record(csat="not a rating", submitted_at="not a date"),
                         source_id=self.db.execute("SELECT source_id FROM source_files WHERE path=?", (BOX.path,)).fetchone()[0], row=9999, end_row=9999)
        insert(db, "source_observations", {"observation_id":record.observation_id,"source_id":record.source_id,
            "sheet":"CSV","source_row":record.row,"source_end_row":record.end_row,"layout_version":record.layout,
            "kind":"feedback","raw_fields_json":raw_json(record.values),"disposition":"pending"})
        event, problems = normalise_event(record)
        insert(db, "feedback_events", event)
        db.execute("UPDATE source_observations SET event_id=?,disposition='retained' WHERE observation_id=?",(event["event_id"],record.observation_id))
        for rule, message, disposition in problems:
            issue(db,record.source_id,rule,message,disposition,record=record)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM quality_issues WHERE observation_id=?", (record.observation_id,)).fetchone()[0],2)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM event_counts WHERE month IS NULL").fetchone()[0],0)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM feedback_events WHERE event_id=? AND rating_valid=0 AND time_eligible=0 AND has_text=1",(event["event_id"],)).fetchone()[0],1)

    def test_foreign_keys_reject_orphaned_references(self):
        db = connect_database(":memory:")
        self.addCleanup(db.close)
        self.db.backup(db)
        with self.assertRaises(sqlite3.IntegrityError):
            db.execute("""INSERT INTO business_monthly VALUES ('MISSING','2026-01','Service',1,100,'MISSING',1,'{}')""")

    def test_demo_queries_execute(self):
        for query in DEMO_QUERIES.values():
            with self.subTest(query=query):
                self.assertGreater(len(self.db.execute(query).fetchall()),0)

    def test_review_report_contains_no_customer_text_or_contact_details(self):
        report=self.output["report"].read_text()
        self.assertFalse(re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}',report))
        self.assertFalse(any(r[0] in report for r in self.db.execute("SELECT comment_raw FROM feedback_events WHERE length(comment_raw)>40")))

    def test_repeat_run_has_identical_content_and_query_results(self):
        again=run_pilot(ROOT,Path(self.temporary.name)/'second')
        second=connect_database(again['database'])
        try:
            self.assertEqual(snapshot(self.db),snapshot(second))
        finally:
            second.close()
        self.assertEqual(self.output['report'].read_bytes(),again['report'].read_bytes())
        self.assertEqual(protected_hashes(ROOT),self.before)

    def test_failed_import_does_not_replace_previous_successful_outputs(self):
        original_database=self.output['database'].read_bytes()
        original_report=self.output['report'].read_bytes()
        with patch('kestrel.pilot.load_feedback',side_effect=PilotError('Synthetic import failure')):
            with self.assertRaisesRegex(PilotError,'Synthetic import failure'):
                run_pilot(ROOT,self.output['database'].parent)
        self.assertEqual(hashlib.sha256(original_database).hexdigest(),hashlib.sha256(self.output['database'].read_bytes()).hexdigest())
        self.assertEqual(original_report,self.output['report'].read_bytes())

    def test_output_directory_cannot_be_a_source_directory(self):
        with self.assertRaisesRegex(PilotError,'supplied inputs'):
            run_pilot(ROOT,ROOT/'feedback')

    def test_output_directory_cannot_be_a_linked_source_directory(self):
        with TemporaryDirectory() as folder:
            base = Path(folder)
            inputs = base / 'inputs'
            inputs.mkdir()
            (inputs / 'reference').mkdir()
            source = base / 'source'
            source.mkdir()
            (inputs / 'feedback').symlink_to(source, target_is_directory=True)
            with self.assertRaisesRegex(PilotError, 'supplied inputs'):
                run_pilot(inputs, source)
            self.assertEqual(list(source.iterdir()), [])

    def test_version_name_alone_is_not_evidence_of_duplication(self):
        def changed_in_memory(root,source):
            extracted=read_source(root,source)
            if source.path.endswith('Chermside Q1 2026 FINAL.xlsx'):
                extracted.records[0].values['Comments']='Synthetic unmatched version marker'
            return extracted
        with patch('kestrel.pilot.read_source',side_effect=changed_in_memory):
            with self.assertRaisesRegex(PilotError,'unique full-row match'):
                run_pilot(ROOT,Path(self.temporary.name)/'unmatched-version')

    def test_bad_aggregate_control_total_blocks_publication(self):
        def changed_in_memory(root,source):
            extracted=read_source(root,source)
            if source.kind=='aggregate':
                extracted.records[0].values['Grand Total']=999
            return extracted
        with patch('kestrel.pilot.read_source',side_effect=changed_in_memory):
            with self.assertRaisesRegex(PilotError,'aggregate counts do not reconcile'):
                run_pilot(ROOT,Path(self.temporary.name)/'invalid-aggregate')


if __name__ == '__main__':
    unittest.main()
